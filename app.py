from io import BytesIO

import json
from pathlib import Path

import streamlit as st #hosting
from openpyxl import load_workbook #manipulating excel files
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

def InsertStandaardObject(workbook):
    json_path = Path(__file__).with_name("standaard_objecten.json")
    with json_path.open(encoding="utf-8") as file:
        objects = json.load(file)

    excel_data_path = Path(__file__).with_name("excel_objecten.json")
    with excel_data_path.open(encoding="utf-8") as file:
        excel_objects = json.load(file)

    object_options = [
        (item["Omschrijving"], item["Nr."])
        for item in objects
        if item.get("Omschrijving") and item.get("Nr.")
    ]
    if not object_options:
        raise ValueError("No complete Omschrijving/Nr. pairs found in standaard_objecten.json")
    allowed_descriptions = {description for description, _ in object_options}
    excel_numbers = {}
    for item in excel_objects:
        description = item.get("Omschrijving")
        number = item.get("Nr.")
        if description in allowed_descriptions and number:
            excel_numbers.setdefault(description, number)
    object_options = [
        (description, excel_numbers.get(description, number))
        for description, number in object_options
    ]

    worksheet = workbook.active
    headers = {
        str(worksheet.cell(row=1, column=column).value).strip().casefold(): column
        for column in range(1, worksheet.max_column + 1)
        if worksheet.cell(row=1, column=column).value is not None
    }
    source_description_column = (
        headers.get("omschrijving")
        or headers.get("object omschrijving")
        or 1
    )
    last_row = max(worksheet.max_row, 2)
    source_descriptions = {
        row: worksheet.cell(row=row, column=source_description_column).value
        for row in range(2, last_row + 1)
    }
    if headers.get("nr.") == 2 and headers.get("omschrijving") == 3:
        number_column = 2
        description_column = 3
    else:
        worksheet.insert_cols(2)
        worksheet.cell(row=1, column=2, value="Nr.")
        worksheet.insert_cols(3)
        worksheet.cell(row=1, column=3, value="Omschrijving")
        number_column = 2
        description_column = 3

    # Create a hidden sheet to store the dropdown options
    options_sheet_name = "_AppOmschrijvingOptions"
    while options_sheet_name in workbook.sheetnames:
        options_sheet_name += "_"
    options_sheet = workbook.create_sheet(options_sheet_name)
    for row, (description, number) in enumerate(object_options, start=1):
        options_sheet.cell(row=row, column=1, value=description)
        options_sheet.cell(row=row, column=2, value=number)
    options_sheet.sheet_state = "hidden"

    defined_name = "AppOmschrijvingOptions"
    while defined_name in workbook.defined_names:
        defined_name += "_"
    workbook.defined_names.add(
        DefinedName(
            defined_name,
            attr_text=f"'{options_sheet_name}'!$A$1:$A${len(object_options)}",
        )
    )

    dropdown = DataValidation(
        type="list",
        formula1=f"={defined_name}",
        allow_blank=False,
        showErrorMessage=True,
        errorStyle="stop",
        errorTitle="Ongeldige invoer",
        error="Kies een geldige Omschrijving uit de lijst.",
    )
    worksheet.add_data_validation(dropdown)
    dropdown.add(f"C2:C{last_row}")

    for row in range(2, last_row + 1):
        worksheet.cell(row=row, column=number_column).value = (
            f'=IFERROR(VLOOKUP(C{row},\'{options_sheet_name}\'!$A:$B,2,FALSE),"")'
        )
        source_description = source_descriptions[row]
        worksheet.cell(row=row, column=description_column).value = (
            source_description
            if source_description in allowed_descriptions
            else None
        )

def MakeThirdColumnRed(workbook):
    worksheet = workbook.active
    red_fill = PatternFill(fill_type="solid", fgColor="FF0000")
    blue_fill = PatternFill(fill_type="solid", fgColor="0000FF")

    for row in range(2, worksheet.max_row + 1):
        if row % 2 == 0:
            worksheet.cell(row=row, column=3).fill = red_fill
        else:
            worksheet.cell(row=row, column=3).fill = blue_fill

st.title("Excel processor")
st.text("Deze app voegt automatisch het nummer en de omschrijving van standaardobjecten toe.")
st.subheader("version 1.010")

uploaded_file = st.file_uploader("Choose an Excel file", type=["xlsx"])

if uploaded_file is not None:
    workbook = load_workbook(uploaded_file)

    # Add your Excel modifications here.
    InsertStandaardObject(workbook)
    # MakeThirdColumnRed(workbook)

    output = BytesIO()
    workbook.save(output)

    st.download_button(
        "Download modified workbook",
        data=output.getvalue(),
        file_name=f"{uploaded_file.name}",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )